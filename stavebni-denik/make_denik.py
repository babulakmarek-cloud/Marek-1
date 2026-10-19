"""Vygeneruje stavebni-denik.xlsx - stavební deník / přehled dokončené práce a montáží."""
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule, CellIsRule
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_rows(ws, first, last, ncols, fmt=None, height=None):
    for r in range(first, last + 1):
        if height:
            ws.row_dimensions[r].height = height
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = WRAP
            if fmt and c in fmt:
                cell.number_format = fmt[c]


wb = Workbook()

# ---------------- Zakázka ----------------
z = wb.active
z.title = "Zakázka"
z.column_dimensions["A"].width = 28
z.column_dimensions["B"].width = 60
z["A1"] = "ZAKÁZKA / ZÁKAZNÍK"
z["A1"].font = Font(bold=True, size=14)
rows = [
    ("Číslo zakázky", "2026-001"),
    ("Název zakázky", "Rekonstrukce topení – ukázka"),
    ("Zákazník", ""),
    ("Kontaktní osoba / telefon / e-mail", ""),
    ("Adresa stavby", ""),
    ("IČO / DIČ zákazníka", ""),
    ("Datum zahájení", date(2026, 10, 19)),
    ("Plánované dokončení", date(2026, 11, 20)),
    ("Cena dle nabídky (Kč bez DPH)", 0),
    ("Poznámky", ""),
]
for i, (k, v) in enumerate(rows, 3):
    z.cell(row=i, column=1, value=k).font = Font(bold=True)
    c = z.cell(row=i, column=2, value=v)
    c.fill = INPUT
    c.border = BORDER
    c.alignment = WRAP
    if isinstance(v, date):
        c.number_format = "DD.MM.YYYY"
    if k.startswith("Cena"):
        c.number_format = '#,##0 "Kč"'
z["A15"] = "Žluté buňky vyplňujete, zelené se počítají samy."
z["A15"].font = Font(italic=True)

# ---------------- Stavební deník ----------------
d = wb.create_sheet("Stavební deník")
cols = ["Datum", "Den", "Počasí", "Teplota °C", "Počet pracovníků", "Odpracováno hodin",
        "Místo / podlaží", "Provedené práce", "Foto 1", "Foto 2", "Popis fotografií",
        "Použitý materiál", "Problémy / překážky", "Dokončeno %", "Zapsal"]
widths = [12, 6, 12, 9, 11, 11, 16, 45, 28, 28, 32, 28, 28, 11, 12]
header(d, 1, cols, widths)
samples = [
    (date(2026, 10, 19), "Zahájení, rozvody v 1.NP", "Slunečno", 14, 2, 16, "1.NP",
     "Demontáž starých rozvodů, osazení nových rozdělovačů.",
     "", "", "Foto 1: stav před. Foto 2: nový rozdělovač.", "Trubka PE-X 16 – 60 m",
     "", 15, "Marek"),
    (date(2026, 10, 20), "", "Oblačno", 11, 2, 16, "1.NP",
     "Pokládka potrubí, tlaková zkouška okruhu 1.",
     "", "", "Foto 1: potrubí před zalitím. Foto 2: manometr.", "Trubka PE-X 16 – 120 m",
     "Zpoždění dodávky fitinek.", 30, "Marek"),
]
for r, s in enumerate(samples, 2):
    for c, v in enumerate(s, 1):
        d.cell(row=r, column=c, value=v)
    d.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"ddd"))')
N = 200
style_rows(d, 2, N, len(cols), fmt={1: "DD.MM.YYYY", 14: "0"}, height=110)
for r in range(2, N + 1):
    d.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"ddd"))').fill = CALC
    d.cell(row=r, column=14).number_format = '0"%"'
dv = DataValidation(type="list", formula1='"Slunečno,Oblačno,Déšť,Sníh,Vítr,Mráz"', allow_blank=True)
d.add_data_validation(dv)
dv.add(f"C2:C{N}")
d.cell(row=1, column=17, value="Foto: vložit → Obrázky → Umístit do buňky (nebo přetáhnout). Řádek má výšku 110 pro náhled.")

# ---------------- Fotodokumentace ----------------
f = wb.create_sheet("Fotodokumentace")
cols = ["Datum", "Místo / podlaží", "Fáze", "Fotografie", "Popis", "Související práce", "Autor"]
header(f, 1, cols, [12, 18, 12, 40, 45, 28, 12])
style_rows(f, 2, 300, len(cols), fmt={1: "DD.MM.YYYY"}, height=150)
dv2 = DataValidation(type="list", formula1='"Před,Během,Po,Závada,Předání"', allow_blank=True)
f.add_data_validation(dv2)
dv2.add("C2:C300")
f["A2"], f["B2"], f["C2"], f["E2"] = date(2026, 10, 19), "1.NP", "Před", "Ukázka: stav před zahájením."

# ---------------- Materiál (přehled) + Spotřeba (log) ----------------
m = wb.create_sheet("Materiál")
cols = ["Položka", "Jednotka", "Plán (množství)", "Spotřebováno", "Zbývá", "Čerpání %",
        "Cena / jedn. (Kč)", "Náklad (Kč)", "Dodavatel"]
header(m, 1, cols, [34, 10, 14, 14, 12, 12, 14, 16, 22])
mat = [("Trubka PE-X 16", "m", 500, 18.5), ("Rozdělovač 8 okruhů", "ks", 4, 2450),
       ("Fitinka 16×2", "ks", 120, 38), ("Otopné těleso 22/600/1000", "ks", 10, 3200)]
MAXM = 60
for r in range(2, MAXM + 1):
    if r - 2 < len(mat):
        n, u, p, c = mat[r - 2]
        m.cell(row=r, column=1, value=n)
        m.cell(row=r, column=2, value=u)
        m.cell(row=r, column=3, value=p)
        m.cell(row=r, column=7, value=c)
    m.cell(row=r, column=4, value=f'=IF(A{r}="","",SUMIF(\'Spotřeba materiálu\'!$B:$B,A{r},\'Spotřeba materiálu\'!$C:$C))')
    m.cell(row=r, column=5, value=f'=IF(A{r}="","",C{r}-D{r})')
    m.cell(row=r, column=6, value=f'=IF(OR(A{r}="",C{r}=0),"",D{r}/C{r})')
    m.cell(row=r, column=8, value=f'=IF(A{r}="","",D{r}*G{r})')
    for c in (4, 5, 6, 8):
        m.cell(row=r, column=c).fill = CALC
    for c in (1, 2, 3, 7, 9):
        m.cell(row=r, column=c).fill = INPUT
style_rows(m, 2, MAXM, len(cols), fmt={6: "0%", 7: "#,##0.00", 8: "#,##0"}, height=None)
m.conditional_formatting.add(f"F2:F{MAXM}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))
m.conditional_formatting.add(f"E2:E{MAXM}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

s = wb.create_sheet("Spotřeba materiálu")
header(s, 1, ["Datum", "Položka", "Množství", "Poznámka"], [12, 34, 12, 40])
style_rows(s, 2, 1000, 4, fmt={1: "DD.MM.YYYY"})
dvm = DataValidation(type="list", formula1=f"=Materiál!$A$2:$A${MAXM}", allow_blank=True)
s.add_data_validation(dvm)
dvm.add("B2:B1000")
for r, (dt, it, q) in enumerate([(date(2026, 10, 19), "Trubka PE-X 16", 60),
                                 (date(2026, 10, 19), "Rozdělovač 8 okruhů", 1),
                                 (date(2026, 10, 20), "Trubka PE-X 16", 120)], 2):
    s.cell(row=r, column=1, value=dt)
    s.cell(row=r, column=2, value=it)
    s.cell(row=r, column=3, value=q)

# ---------------- Pokrok ----------------
p = wb.create_sheet("Pokrok")
cols = ["Etapa", "Váha", "Plán od", "Plán do", "Skutečně od", "Skutečně do", "Hotovo %", "Stav", "Poznámka"]
header(p, 1, cols, [34, 8, 12, 12, 12, 12, 11, 14, 30])
base = date(2026, 10, 19)
et = [("Demontáž stávajícího", 10, 0, 2, 100), ("Hrubé rozvody", 30, 2, 10, 60),
      ("Tlakové zkoušky", 10, 10, 12, 0), ("Osazení těles a armatur", 25, 12, 18, 0),
      ("Napojení zdroje, zprovoznění", 15, 18, 22, 0), ("Předání zákazníkovi", 10, 22, 24, 0)]
MAXP = 30
for r in range(2, MAXP + 1):
    if r - 2 < len(et):
        n, w, a, b, h = et[r - 2]
        p.cell(row=r, column=1, value=n)
        p.cell(row=r, column=2, value=w)
        p.cell(row=r, column=3, value=base + timedelta(days=a))
        p.cell(row=r, column=4, value=base + timedelta(days=b))
        p.cell(row=r, column=7, value=h / 100)
    p.cell(row=r, column=8, value=f'=IF(A{r}="","",IF(G{r}>=1,"Hotovo",IF(G{r}>0,"Rozpracováno","Čeká")))')
    p.cell(row=r, column=8).fill = CALC
style_rows(p, 2, MAXP, len(cols), fmt={3: "DD.MM.YYYY", 4: "DD.MM.YYYY", 5: "DD.MM.YYYY", 6: "DD.MM.YYYY", 7: "0%"})
p.conditional_formatting.add(f"G2:G{MAXP}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))
dvp = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
p.add_data_validation(dvp)
dvp.add(f"G2:G{MAXP}")

# ---------------- Přehled ----------------
o = wb.create_sheet("Přehled", 0)
o.column_dimensions["A"].width = 38
o.column_dimensions["B"].width = 22
o["A1"] = "PŘEHLED ZAKÁZKY"
o["A1"].font = Font(bold=True, size=16)
kpi = [
    ("Zakázka", "=Zakázka!B4", None),
    ("Zákazník", "=Zakázka!B5&\"\"", None),
    ("Celkový pokrok (vážený)", f'=IFERROR(SUMPRODUCT(Pokrok!B2:B{MAXP},Pokrok!G2:G{MAXP})/SUM(Pokrok!B2:B{MAXP}),0)', "0%"),
    ("Počet zápisů v deníku", "=COUNT('Stavební deník'!A2:A200)", "0"),
    ("Odpracováno hodin celkem", "=SUM('Stavební deník'!F2:F200)", "#,##0"),
    ("Počet fotografií v deníku (řádky s popisem)", "=COUNTA('Stavební deník'!K2:K200)", "0"),
    ("Náklad na materiál (Kč)", f"=SUM(Materiál!H2:H{MAXM})", "#,##0"),
    ("Cena dle nabídky (Kč)", "=Zakázka!B11", "#,##0"),
    ("Hrubá marže před prací (Kč)", "=B10-B9", "#,##0"),
    ("Dnů od zahájení", '=IF(Zakázka!B9="","",TODAY()-Zakázka!B9)', "0"),
    ("Dnů do plánovaného dokončení", '=IF(Zakázka!B10="","",Zakázka!B10-TODAY())', "0"),
]
for i, (k, v, fm) in enumerate(kpi, 3):
    a = o.cell(row=i, column=1, value=k)
    a.font = Font(bold=True)
    a.border = BORDER
    b = o.cell(row=i, column=2, value=v)
    b.fill = CALC
    b.border = BORDER
    if fm:
        b.number_format = fm
o.conditional_formatting.add("B5", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))

wb.save("/home/user/Marek-1/stavebni-denik/stavebni-denik.xlsx")
print("ok")
