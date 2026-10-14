"""Pomocné třídy pro generování Excelu (styl shodný s elektrikar/)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

INP = PatternFill('solid', fgColor='FFF2CC')
OUTF = PatternFill('solid', fgColor='E8F0FC')
HEAD = PatternFill('solid', fgColor='0B63D1')
SUB = PatternFill('solid', fgColor='D9E0E8')
thin = Side(style='thin', color='B8C2CE')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
GOOD = PatternFill('solid', fgColor='C6EFCE')
BADF = PatternFill('solid', fgColor='FFC7CE')

wb = Workbook()
wb.remove(wb.active)


class Sheet:
    def __init__(self, title, intro='', widths=(42, 16, 16)):
        self.ws = wb.create_sheet(title)
        self.r = 1
        ws = self.ws
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws['A1'] = title
        ws['A1'].font = Font(size=15, bold=True, color='0B63D1')
        self.r = 2
        if intro:
            ws.cell(self.r, 1, intro).font = Font(italic=True, color='5B6877')
            ws.cell(self.r, 1).alignment = Alignment(wrap_text=False)
            self.r += 1
        self.r += 1

    def head(self, text):
        c = self.ws.cell(self.r, 1, text)
        for col in (1, 2, 3):
            self.ws.cell(self.r, col).fill = HEAD
        c.font = Font(bold=True, color='FFFFFF')
        self.r += 1

    def inp(self, label, val, unit='', options=None, fmt=None):
        ws = self.ws
        ws.cell(self.r, 1, label)
        c = ws.cell(self.r, 2, val)
        c.fill = INP
        c.border = BOX
        c.font = Font(color='0000FF', bold=True)
        if fmt:
            c.number_format = fmt
        ws.cell(self.r, 3, unit).font = Font(color='5B6877')
        if options:
            dv = DataValidation(type='list', formula1='"' + ','.join(str(o) for o in options) + '"', allow_blank=False)
            ws.add_data_validation(dv)
            dv.add(c)
        addr = f'$B${self.r}'
        self.r += 1
        return addr

    def calc(self, label, formula, unit='', fmt='0.00', bold=False):
        ws = self.ws
        ws.cell(self.r, 1, label)
        c = ws.cell(self.r, 2, formula)
        c.fill = OUTF
        c.border = BOX
        c.number_format = fmt
        c.font = Font(bold=bold)
        c.alignment = Alignment(horizontal='right')
        ws.cell(self.r, 3, unit).font = Font(color='5B6877')
        addr = f'$B${self.r}'
        self.r += 1
        return addr

    def note(self, text):
        c = self.ws.cell(self.r, 1, text)
        c.font = Font(italic=True, size=9, color='5B6877')
        self.r += 1

    def gap(self, n=1):
        self.r += n

    def status_cf(self, rng):
        self.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("OK",{rng.split(":")[0]}))'], fill=GOOD))
        self.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("CHYBA",{rng.split(":")[0]}))'], fill=BADF))
        self.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("POZOR",{rng.split(":")[0]}))'], fill=PatternFill('solid', fgColor='FFEB9C')))


def table(ws, r0, c0, header, rows, widths=None, numfmt=None):
    for j, h in enumerate(header):
        c = ws.cell(r0, c0 + j, h)
        c.font = Font(bold=True)
        c.fill = SUB
        c.border = BOX
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(r0 + 1 + i, c0 + j, v)
            c.border = BOX
            if numfmt and j in numfmt:
                c.number_format = numfmt[j]


