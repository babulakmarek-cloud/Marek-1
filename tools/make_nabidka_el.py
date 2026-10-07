#!/usr/bin/env python3
"""Cenová nabídka: elektro přívody pro klimatizaci (2x venkovní jednotka Haier 2+1).

Vytvoří vystupy/Nabidka_elektro_privody_klimatizace.xlsx (ceny jsou v buňkách
upravitelné, součty jsou vzorce). PDF se z něj dělá přes LibreOffice.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

OUT = "vystupy/Nabidka_elektro_privody_klimatizace.xlsx"

thin = Side(style="thin", color="000000")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
GREY = PatternFill("solid", fgColor="D9D9D9")
LIGHT = PatternFill("solid", fgColor="F2F2F2")
KC = '#,##0.00" Kč"'

wb = Workbook()
ws = wb.active
ws.title = "Nabídka"
for col, w in zip("ABCDEF", (8, 58, 9, 9, 16, 18)):
    ws.column_dimensions[col].width = w

r = 1
ws.cell(r, 1, "Cenová nabídka č. 26_960-EL").font = Font(bold=True, size=16)
r += 1
ws.cell(r, 1, "ze dne: 07.10.2026").font = Font(size=11)
r += 1
ws.cell(r, 1, "Elektro přívody pro klimatizaci RD – 2× venkovní jednotka "
              "HAIER Multi-Split 2+1 (soupis č. 26_960, Cool Technology s.r.o.)")
ws.cell(r, 1).font = Font(bold=True, size=11)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r].height = 32
r += 2

# --- adresy
ws.cell(r, 1, "Dodavatel:").font = Font(bold=True)
ws.cell(r, 4, "Odběratel / místo realizace:").font = Font(bold=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
r += 1
supplier = ["[doplnit jméno / obchodní firmu]", "[doplnit adresu]",
            "IČO: [doplnit]   DIČ: [doplnit]", "Tel.: [doplnit]   E-mail: babulakmarek@gmail.com"]
customer = ["pan Pavel Znojemský", "Lesná 33", "405 02 Děčín", ""]
for s, c in zip(supplier, customer):
    ws.cell(r, 1, s)
    ws.cell(r, 4, c)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    r += 1
r += 1

# --- hlavička tabulky
heads = ["Č.", "Název", "Jedn.", "Počet", "Jedn. cena\nbez DPH", "Cena celkem\nbez DPH"]
for i, h in enumerate(heads, 1):
    c = ws.cell(r, i, h)
    c.font = Font(bold=True)
    c.fill = GREY
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[r].height = 32
header_row = r
r += 1

num = 0
first_item = r
item_rows = []


def section(title):
    global r
    for i in range(1, 7):
        ws.cell(r, i).border = BORDER
        ws.cell(r, i).fill = LIGHT
    ws.cell(r, 2, title).font = Font(bold=True)
    ws.cell(r, 2).alignment = Alignment(horizontal="center")
    r += 1


def item(name, unit, qty, price, bold=False):
    global r, num
    num += 1
    ws.cell(r, 1, f"{num:02d}")
    ws.cell(r, 2, name)
    ws.cell(r, 3, unit)
    ws.cell(r, 4, qty)
    ws.cell(r, 5, price)
    ws.cell(r, 6, f"=D{r}*E{r}")
    for i in range(1, 7):
        c = ws.cell(r, i)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=(i == 2),
                                horizontal="center" if i in (1, 3, 4) else None)
        if bold:
            c.font = Font(bold=True)
    ws.cell(r, 5).number_format = KC
    ws.cell(r, 6).number_format = KC
    ws.cell(r, 4).number_format = "0.##"
    item_rows.append(r)
    r += 1


def feed(title, room_note):
    section(title)
    item(f"Kabel CYKY-J 3×2,5 mm² – napájení venkovní jednotky ({room_note}), "
         "délka trasy cca 10 m", "m", 10, 45)
    item("Plastová elektroinstalační lišta bílá 40×25 mm vč. tvarovek "
         "(rohy, spojky, koncovky), cca 10 m", "bm", 10, 70)
    item("Jistič 1pólový B16A, charakteristika B, 6 kA", "ks", 1, 180)
    item("Drobný instalační materiál (svorky, dutinky, hmoždinky, "
         "popisné štítky)", "soubor", 1, 250)
    item("Montáž lišty a uložení kabelu do lišty", "bm", 10, 150)
    item("Montáž jističe do stávajícího rozvaděče, připojení a popis okruhu",
         "ks", 1, 450)
    item("Připojení kabelu na svorkovnici venkovní jednotky, kontrola "
         "zapojení", "kpl", 1, 350)


feed("Přívod č. 1 – venkovní jednotka 2+1 (kuchyň + ložnice)", "jednotka č. 1")
feed("Přívod č. 2 – venkovní jednotka 2+1 (dětský pokoj č. 1 + č. 2)", "jednotka č. 2")

section("Společné položky")
item("Výchozí revize elektrické instalace nových přívodů (měření, "
     "protokol, revizní zpráva)", "kpl", 1, 2500)
item("Doprava materiálu a pracovníků – paušál", "pauš", 1, 400)
last_item = r - 1

# --- součty
sum_rng = f"F{first_item}:F{last_item}"
for label, formula, bold in (
    ("Cena celkem bez DPH:", f"=SUM({sum_rng})", True),
    ("DPH 21 %:", None, False),
    ("Cena celkem vč. DPH:", None, True),
):
    ws.cell(r, 2, label).font = Font(bold=bold)
    ws.cell(r, 2).alignment = Alignment(horizontal="right")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    for i in range(1, 7):
        ws.cell(r, i).border = BORDER
    if label.startswith("Cena celkem bez"):
        base_row = r
        ws.cell(r, 6, formula)
    elif label.startswith("DPH"):
        ws.cell(r, 6, f"=F{base_row}*0.21")
    else:
        ws.cell(r, 6, f"=F{base_row}+F{base_row + 1}")
    ws.cell(r, 6).number_format = KC
    ws.cell(r, 6).font = Font(bold=bold)
    r += 1
r += 1

# --- nepovinné položky
ws.cell(r, 1, "Nepovinné položky (nejsou zahrnuty v ceně výše – dle výsledku "
              "prohlídky rozvaděče)").font = Font(bold=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
r += 1
opt_heads = ["", "Název", "Jedn.", "Počet", "Jedn. cena\nbez DPH", "Cena celkem\nbez DPH"]
for i, h in enumerate(opt_heads, 1):
    c = ws.cell(r, i, h)
    c.font = Font(bold=True)
    c.fill = GREY
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[r].height = 32
r += 1
options = [
    ("A", "Proudový chránič s nadproudovou ochranou RCBO B16A/30 mA (1P+N) "
          "místo samotného jističe B16A – příplatek za 1 ks vč. montáže", "ks", 2, 850),
    ("B", "Servisní vypínač 16A IP65 u venkovní jednotky (vč. krabice a "
          "zapojení)", "ks", 2, 650),
    ("C", "Plastová rozvodnice 1×8 modulů (pokud ve stávajícím rozvaděči "
          "není místo pro 2 jističe) vč. montáže a propojení", "ks", 1, 1800),
]
for code, name, unit, qty, price in options:
    ws.cell(r, 1, code)
    ws.cell(r, 2, name)
    ws.cell(r, 3, unit)
    ws.cell(r, 4, qty)
    ws.cell(r, 5, price)
    ws.cell(r, 6, f"=D{r}*E{r}")
    for i in range(1, 7):
        c = ws.cell(r, i)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=(i == 2),
                                horizontal="center" if i in (1, 3, 4) else None)
    ws.cell(r, 5).number_format = KC
    ws.cell(r, 6).number_format = KC
    ws.row_dimensions[r].height = 44
    r += 1
r += 1

# --- poznámky
ws.cell(r, 1, "Poznámky a předpoklady:").font = Font(bold=True, size=11)
r += 1
notes = [
    "Nabídka je platná 30 dní od data vystavení.",
    "Napájení každé venkovní jednotky (Haier 2U50S2SM1FA-3, 1~ 230 V/50 Hz) je "
    "samostatný okruh jištěný 1fázovým jističem B16A, kabel CYKY-J 3×2,5 mm² "
    "(dle technické dokumentace výrobce: napájecí kabel 3×2,5 mm², max. provozní "
    "proud 9,0 A).",
    "Kabeláž je vedena v plastové liště; délka každé trasy cca 10 m. Nabídka je "
    "kalkulována na délku 2× 10 m, materiál se vyúčtuje dle skutečného provedení.",
    "Předpokládá se volná kapacita hlavního jištění a volné místo ve stávajícím "
    "rozvaděči pro 2 jističe; případně viz nepovinná položka C. Fázové "
    "rozložení zátěže (L1/L2/L3) bude vyváženo při montáži.",
    "Nabídka zahrnuje výchozí revizi nově zřízených přívodů a revizní zprávu "
    "(v soupisu č. 26_960 firmy Cool Technology s.r.o. není revize el. přívodu "
    "ani jeho montáž obsažena).",
    "Nabídka neobsahuje: komunikační vedení a chladivové potrubí mezi vnitřními "
    "a venkovními jednotkami, prostupy obvodovými a dělicími stěnami, zednické "
    "a malířské práce, montáž ani zprovoznění klimatizačních jednotek "
    "(zajišťuje Cool Technology s.r.o.).",
    "Ceny jsou uvedeny bez DPH; DPH ve výši 21 % je vyčísleno v součtu.",
    "Nabídka vychází z fotodokumentace stávajícího rozvaděče a elektroměrového "
    "rozvaděče; konečná cena bude upřesněna po prohlídce na místě.",
]
for n in notes:
    ws.cell(r, 1, "•")
    ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(r, 2, n)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(r, 2).font = Font(size=10)
    ws.row_dimensions[r].height = 15 * max(1, -(-len(n) // 82))
    r += 1
r += 1
ws.cell(r, 1, "Vystavil:")
ws.cell(r, 2, "[doplnit jméno]")
r += 1
ws.cell(r, 1, "Tel.:")
ws.cell(r, 2, "[doplnit telefon]")

# --- tisk
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)

wb.save(OUT)
print("saved", OUT)
