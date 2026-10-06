import sys, json, re, math, subprocess, openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
texts = json.load(open(sys.argv[1])); pipes_txt = open(sys.argv[2]).read().splitlines()[1:]
out = sys.argv[3]
wb = openpyxl.Workbook()
def sheet(name, head, rows):
    ws = wb.create_sheet(name); ws.append(head)
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows: ws.append(r)
    for i,_ in enumerate(head,1): ws.column_dimensions[get_column_letter(i)].width = 28
    ws.freeze_panes = "A2"
wb.remove(wb.active)
# potrubí
rows=[]
for l in pipes_txt:
    m = re.match(r"(.+?)\s+(\d+)\s+([\d.]+)$", l.strip())
    if m:
        name=m.group(1); temp=re.search(r"\+(\d+)",name)
        kind = "přívod" if re.search(r"\bV\b|Vorlauf",name) else ("zpátečka" if re.search(r"\bR\b|Rücklauf",name) else "?")
        rows.append([name, int(temp.group(1)) if temp else None, kind, int(m.group(2)), float(m.group(3))])
sheet("Potrubí dle hladin",["Hladina","Teplota °C","Přívod/zpátečka","Počet prvků","Délka m (odhad z geometrie)"],rows)
# otopná tělesa
body=[t for t in texts if re.search(r"heizk",t["t"],re.I)]
pw=[t for t in texts if re.match(r"P\s*[≈=]\s*[\d.,]+\s*W",t["t"])]
used=set(); rows=[]
for b in body:
    best=None
    for i,p in enumerate(pw):
        if i in used: continue
        dd=math.hypot(b["x"]-p["x"],b["y"]-p["y"])
        if best is None or dd<best[0]: best=(dd,i)
    pwr=None; dist=None
    if best and best[0]<6000:
        used.add(best[1]); pwr=int(re.sub(r"[^\d]","",re.sub(r"[.,]\d+(?=\s*W)","",pw[best[1]]["t"].split("W")[0]))); dist=round(best[0])
    n=re.match(r"(\d+)\s*x\s*",b["t"]); cnt=int(n.group(1)) if n else 1
    rows.append([b["t"],cnt,pwr,dist,b["x"],b["y"],b["layer"]])
sheet("Otopná tělesa",["Popisek","Počet","Výkon W (nejbližší popisek P≈)","Vzdálenost popisků mm (kontrola párování)","X mm","Y mm","Hladina"],rows)
# DN
rows=[[t["t"],int(re.sub(r"\D","",t["t"])),t["x"],t["y"],t["layer"]] for t in texts if re.fullmatch(r"DN\s?\d+",t["t"])]
sheet("DN popisky",["Popisek","DN","X mm","Y mm","Hladina"],rows)
# místnosti a ostatní popisy
rows=[[t["t"],t["x"],t["y"],t["layer"]] for t in texts if t["layer"] in ("A06_Text","A06_Text  -Belegung","A0_Grundriß") and not re.fullmatch(r"[\d.,\s]+",t["t"])]
sheet("Popisy dispozice",["Text","X mm","Y mm","Hladina"],rows)
wb.save(out); print("saved",out,[ (ws.title,ws.max_row-1) for ws in wb])
